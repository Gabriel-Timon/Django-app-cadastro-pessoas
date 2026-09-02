import json
from datetime import date, timedelta
from io import BytesIO, StringIO

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .forms import PessoaForm
from .models import Pessoa


class SecureRequestMixin:
    """Executa requests como HTTPS, inclusive quando o CI usa settings de produção."""

    def setUp(self):
        super().setUp()
        self.client.defaults['HTTP_X_FORWARDED_PROTO'] = 'https'


def cpf_from_base(base: str) -> str:
    digits = [int(digit) for digit in base]
    for tamanho in (9, 10):
        soma = sum(digit * weight for digit, weight in zip(digits[:tamanho], range(tamanho + 1, 1, -1)))
        verifier = (soma * 10) % 11
        digits.append(0 if verifier == 10 else verifier)
    return ''.join(str(digit) for digit in digits)


class PessoaModelTests(TestCase):
    def test_properties_mask_cpf_translate_sexo_and_calculate_age(self):
        hoje = date.today()
        pessoa = Pessoa(
            nome='Ana',
            sobrenome='Silva',
            data_nascimento=date(hoje.year - 30, hoje.month, hoje.day),
            sexo=Pessoa.SEXO_FEM,
            cpf=cpf_from_base('123456789'),
        )

        self.assertEqual(pessoa.cpf_formatado, '123.456.789-09')
        self.assertEqual(pessoa.cpf_mascarado, '***.***.***-09')
        self.assertEqual(pessoa.sexo_extenso, 'Feminino')
        self.assertEqual(pessoa.idade, 30)
        self.assertEqual(str(pessoa), 'Ana Silva (***.***.***-09)')

    def test_model_rejects_invalid_cpf_and_future_birth_date(self):
        pessoa = Pessoa(
            nome='Ana',
            sobrenome='Silva',
            data_nascimento=date.today() + timedelta(days=1),
            sexo=Pessoa.SEXO_FEM,
            cpf='11111111111',
        )

        with self.assertRaises(ValidationError) as error:
            pessoa.full_clean()

        self.assertIn('cpf', error.exception.message_dict)
        self.assertIn('data_nascimento', error.exception.message_dict)


class PessoaFormTests(TestCase):
    def test_form_normalizes_names_and_formatted_cpf(self):
        cpf = cpf_from_base('123456789')
        form = PessoaForm(
            data={
                'nome': '  ana  ',
                'sobrenome': ' silva ',
                'data_nascimento': '10/02/1990',
                'sexo': Pessoa.SEXO_FEM,
                'cpf': f'{cpf[:3]}.{cpf[3:6]}.{cpf[6:9]}-{cpf[9:]}',
            }
        )

        self.assertTrue(form.is_valid(), form.errors)
        pessoa = form.save()

        self.assertEqual(pessoa.nome, 'Ana')
        self.assertEqual(pessoa.sobrenome, 'Silva')
        self.assertEqual(pessoa.cpf, cpf)

    def test_form_rejects_invalid_cpf_empty_name_and_formula_prefix(self):
        form = PessoaForm(
            data={
                'nome': '=HYPERLINK("https://example.com")',
                'sobrenome': 'Silva',
                'data_nascimento': '10/02/1990',
                'sexo': Pessoa.SEXO_FEM,
                'cpf': '123',
            }
        )

        self.assertFalse(form.is_valid())
        self.assertIn('nome', form.errors)
        self.assertIn('cpf', form.errors)


class AccessControlTests(SecureRequestMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username='sem-permissao', password='senha-forte-teste')

    def test_login_and_admin_login_are_public(self):
        self.assertEqual(self.client.get(reverse('login')).status_code, 200)
        self.assertEqual(self.client.get('/admin/login/').status_code, 200)

    def test_anonymous_user_is_redirected_to_login(self):
        response = self.client.get(reverse('pessoas:dashboard'))

        self.assertRedirects(response, f"{reverse('login')}?next={reverse('pessoas:dashboard')}")

    def test_authenticated_user_without_permission_receives_403(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse('pessoas:dashboard'))

        self.assertEqual(response.status_code, 403)

    def test_logout_requires_post_and_ends_session(self):
        self.client.force_login(self.user)

        get_response = self.client.get(reverse('logout'))
        post_response = self.client.post(reverse('logout'))

        self.assertEqual(get_response.status_code, 405)
        self.assertRedirects(post_response, reverse('login'))


class SignupTests(SecureRequestMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.demo_pessoa = Pessoa.objects.create(
            demo_key='portfolio-demo-001',
            nome='Pessoa',
            sobrenome='Demonstração',
            data_nascimento=date(1995, 6, 15),
            sexo=Pessoa.SEXO_FEM,
            cpf=cpf_from_base('456123789'),
        )
        cls.private_pessoa = Pessoa.objects.create(
            nome='Registro',
            sobrenome='Não demonstrativo',
            data_nascimento=date(1990, 1, 1),
            sexo=Pessoa.SEXO_MASC,
            cpf=cpf_from_base('741852963'),
        )
        cls.noncanonical_demo = Pessoa.objects.create(
            demo_key='portfolio-demo-extra',
            nome='Marcado',
            sobrenome='Fora da lista canônica',
            data_nascimento=date(1988, 2, 2),
            sexo=Pessoa.SEXO_MASC,
            cpf=cpf_from_base('852741963'),
        )

    def signup_data(self, **overrides):
        data = {
            'username': 'visitante-demo',
            'email': 'visitante@example.com',
            'password1': 'Senha-forte-de-portfolio-2026!',
            'password2': 'Senha-forte-de-portfolio-2026!',
            'website': '',
        }
        data.update(overrides)
        return data

    def test_signup_is_public_and_creates_logged_in_read_only_user(self):
        self.assertEqual(self.client.get(reverse('pessoas:signup')).status_code, 200)

        response = self.client.post(reverse('pessoas:signup'), self.signup_data())

        self.assertRedirects(response, reverse('pessoas:dashboard'))
        user = get_user_model().objects.get(username='visitante-demo')
        self.assertFalse(user.is_staff)
        self.assertFalse(user.is_superuser)
        self.assertEqual(user.get_all_permissions(), {'pessoas.view_pessoa'})
        self.assertEqual(int(self.client.session['_auth_user_id']), user.pk)

    def test_read_only_user_sees_only_demo_data_and_not_privileged_actions(self):
        self.client.post(reverse('pessoas:signup'), self.signup_data())

        list_response = self.client.get(reverse('pessoas:pessoa_list'))
        dashboard_response = self.client.get(reverse('pessoas:dashboard'))

        self.assertContains(list_response, self.demo_pessoa.nome)
        self.assertNotContains(list_response, self.private_pessoa.sobrenome)
        self.assertNotContains(list_response, self.noncanonical_demo.sobrenome)
        self.assertContains(dashboard_response, 'somente leitura')
        self.assertContains(dashboard_response, 'Menu de filtros')
        self.assertContains(dashboard_response, 'Exibir estat')
        self.assertNotContains(dashboard_response, 'Adicionar pessoa')
        self.assertNotContains(dashboard_response, 'Editar cadastro')
        self.assertNotContains(dashboard_response, 'Opções de exportação')
        self.assertEqual(
            self.client.get(
                reverse('pessoas:pessoa_detail', kwargs={'pk': self.private_pessoa.pk})
            ).status_code,
            404,
        )

    def test_read_only_user_can_query_filters_and_dashboards(self):
        self.client.post(reverse('pessoas:signup'), self.signup_data())

        readable_routes = (
            'pessoas:dashboard',
            'pessoas:pessoa_list',
            'pessoas:menu_filtros',
            'pessoas:filtro_mulheres',
            'pessoas:menu_estatisticas',
            'pessoas:est_homens_mulheres',
            'pessoas:est_faixa_etaria',
        )
        for route in readable_routes:
            with self.subTest(route=route):
                self.assertEqual(self.client.get(reverse(route)).status_code, 200)

    def test_read_only_user_cannot_mutate_export_or_access_admin(self):
        self.client.post(reverse('pessoas:signup'), self.signup_data())

        denied_routes = (
            reverse('pessoas:pessoa_create'),
            reverse('pessoas:pessoa_edicao_menu', kwargs={'pk': self.demo_pessoa.pk}),
            reverse('pessoas:pessoa_delete', kwargs={'pk': self.demo_pessoa.pk}),
            reverse('pessoas:buscar_por_cpf'),
            reverse('pessoas:export_csv'),
        )
        for url in denied_routes:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 403)

        original_name = self.demo_pessoa.nome
        mutation_attempts = (
            (
                reverse('pessoas:pessoa_create'),
                {
                    'nome': 'Invasor',
                    'sobrenome': 'Bloqueado',
                    'data_nascimento': '01/01/2000',
                    'sexo': Pessoa.SEXO_MASC,
                    'cpf': cpf_from_base('159357486'),
                },
            ),
            (
                reverse('pessoas:pessoa_edit_nome', kwargs={'pk': self.demo_pessoa.pk}),
                {'nome': 'Alterado'},
            ),
            (
                reverse('pessoas:pessoa_delete', kwargs={'pk': self.demo_pessoa.pk}),
                {},
            ),
        )
        for url, payload in mutation_attempts:
            with self.subTest(method='POST', url=url):
                self.assertEqual(self.client.post(url, payload).status_code, 403)

        self.demo_pessoa.refresh_from_db()
        self.assertEqual(self.demo_pessoa.nome, original_name)
        self.assertFalse(Pessoa.objects.filter(nome='Invasor').exists())

        admin_response = self.client.get('/admin/')
        self.assertEqual(admin_response.status_code, 302)
        self.assertIn('/admin/login/', admin_response.url)

    def test_signup_ignores_privilege_fields_and_rejects_duplicate_email(self):
        first_data = self.signup_data(is_staff='on', is_superuser='on', user_permissions='1')
        self.client.post(reverse('pessoas:signup'), first_data)
        user = get_user_model().objects.get(username='visitante-demo')
        self.assertFalse(user.is_staff)
        self.assertFalse(user.is_superuser)
        self.assertEqual(user.get_all_permissions(), {'pessoas.view_pessoa'})

        self.client.logout()
        duplicate = self.signup_data(
            username='outro-visitante',
            email='VISITANTE@example.com',
        )
        response = self.client.post(reverse('pessoas:signup'), duplicate)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(get_user_model().objects.filter(username='outro-visitante').exists())

    def test_signup_rejects_honeypot_and_rate_limits_attempts(self):
        honeypot_response = self.client.post(
            reverse('pessoas:signup'),
            self.signup_data(website='https://spam.example'),
            REMOTE_ADDR='198.51.100.10',
        )
        self.assertEqual(honeypot_response.status_code, 200)
        self.assertFalse(get_user_model().objects.filter(username='visitante-demo').exists())

        for _ in range(5):
            response = self.client.post(
                reverse('pessoas:signup'),
                {'username': ''},
                REMOTE_ADDR='203.0.113.10',
            )
            self.assertEqual(response.status_code, 200)

        limited_response = self.client.post(
            reverse('pessoas:signup'),
            {'username': ''},
            REMOTE_ADDR='203.0.113.10',
        )
        self.assertEqual(limited_response.status_code, 429)
        self.assertIn('Retry-After', limited_response)
        self.assertContains(limited_response, 'aproximadamente 60 minutos', status_code=429)
        self.assertNotContains(limited_response, 'Criar conta e entrar', status_code=429)


class DemoSeedTests(TestCase):
    def test_seed_creates_exactly_50_balanced_idempotent_records(self):
        first_output = StringIO()
        call_command('seed_demo_data', stdout=first_output)

        demo_people = Pessoa.objects.filter(demo_key__isnull=False)
        self.assertEqual(demo_people.count(), 50)
        self.assertEqual(demo_people.filter(sexo=Pessoa.SEXO_FEM).count(), 25)
        self.assertEqual(demo_people.filter(sexo=Pessoa.SEXO_MASC).count(), 25)
        self.assertEqual(demo_people.values('demo_key').distinct().count(), 50)
        self.assertIn('50 registros', first_output.getvalue())

        for gender in (Pessoa.SEXO_FEM, Pessoa.SEXO_MASC):
            ages = [person.idade for person in demo_people.filter(sexo=gender)]
            self.assertTrue(any(age <= 17 for age in ages))
            self.assertTrue(any(18 <= age <= 29 for age in ages))
            self.assertTrue(any(30 <= age <= 49 for age in ages))
            self.assertTrue(any(age >= 50 for age in ages))
        for person in demo_people:
            person.full_clean()

        unchanged_timestamp = demo_people.get(
            demo_key='portfolio-demo-001'
        ).atualizado_em
        non_demo = Pessoa.objects.create(
            nome='Registro',
            sobrenome='Preservado',
            data_nascimento=date(1980, 1, 1),
            sexo=Pessoa.SEXO_MASC,
            cpf=cpf_from_base('321654987'),
        )
        second_output = StringIO()
        call_command('seed_demo_data', stdout=second_output)

        self.assertEqual(Pessoa.objects.filter(demo_key__isnull=False).count(), 50)
        self.assertTrue(Pessoa.objects.filter(pk=non_demo.pk).exists())
        self.assertEqual(
            Pessoa.objects.get(demo_key='portfolio-demo-001').atualizado_em,
            unchanged_timestamp,
        )
        self.assertIn('0 criados, 0 atualizados, 50 inalterados', second_output.getvalue())


class PessoaViewTests(SecureRequestMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_superuser(
            username='admin-teste',
            email='admin@example.com',
            password='senha-forte-teste',
        )
        cls.pessoa = Pessoa.objects.create(
            nome='Ana',
            sobrenome='Silva',
            data_nascimento=date(1990, 1, 10),
            sexo=Pessoa.SEXO_FEM,
            cpf=cpf_from_base('123456789'),
        )
        cls.outro = Pessoa.objects.create(
            nome='Bruno',
            sobrenome='Santos',
            data_nascimento=date(1985, 5, 20),
            sexo=Pessoa.SEXO_MASC,
            cpf=cpf_from_base('987654321'),
        )

    def setUp(self):
        super().setUp()
        self.client.force_login(self.user)

    def test_dashboard_displays_summary_and_security_headers(self):
        response = self.client.get(reverse('pessoas:dashboard'))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['total'], 2)
        self.assertEqual(response.context['homens'], 1)
        self.assertEqual(response.context['mulheres'], 1)
        self.assertContains(response, 'Menu Principal')
        self.assertIn("default-src 'self'", response.headers['Content-Security-Policy'])
        self.assertEqual(response.headers['X-Robots-Tag'], 'noindex, nofollow')
        self.assertIn('no-store', response.headers['Cache-Control'])

    def test_list_searches_name_and_masks_cpf(self):
        response = self.client.get(reverse('pessoas:pessoa_list'), {'q': 'Ana'})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Ana')
        self.assertNotContains(response, 'Bruno')
        self.assertContains(response, self.pessoa.cpf_mascarado)
        self.assertNotContains(response, self.pessoa.cpf)

    def test_create_view_persists_person_and_redirects_to_detail(self):
        cpf = cpf_from_base('111222333')
        response = self.client.post(
            reverse('pessoas:pessoa_create'),
            {
                'nome': 'carla',
                'sobrenome': 'souza',
                'data_nascimento': '01/01/2000',
                'sexo': Pessoa.SEXO_FEM,
                'cpf': cpf,
            },
        )

        pessoa = Pessoa.objects.get(cpf=cpf)
        self.assertRedirects(response, reverse('pessoas:pessoa_detail', kwargs={'pk': pessoa.pk}))
        self.assertEqual(pessoa.nome, 'Carla')

    def test_edit_name_view_updates_person(self):
        response = self.client.post(
            reverse('pessoas:pessoa_edit_nome', kwargs={'pk': self.pessoa.pk}),
            {'nome': 'beatriz'},
        )

        self.assertRedirects(response, reverse('pessoas:pessoa_detail', kwargs={'pk': self.pessoa.pk}))
        self.pessoa.refresh_from_db()
        self.assertEqual(self.pessoa.nome, 'Beatriz')

    def test_search_by_cpf_redirects_to_person_detail(self):
        response = self.client.post(reverse('pessoas:buscar_por_cpf'), {'cpf': self.pessoa.cpf_formatado})

        self.assertRedirects(response, reverse('pessoas:pessoa_detail', kwargs={'pk': self.pessoa.pk}))

    def test_delete_view_removes_person(self):
        response = self.client.post(reverse('pessoas:pessoa_delete', kwargs={'pk': self.outro.pk}))

        self.assertRedirects(response, reverse('pessoas:pessoa_list'))
        self.assertFalse(Pessoa.objects.filter(pk=self.outro.pk).exists())

    def test_filter_and_statistics_views(self):
        response = self.client.get(reverse('pessoas:filtro_homens'))
        self.assertContains(response, 'Bruno')
        self.assertNotContains(response, 'Ana')

        response = self.client.get(reverse('pessoas:est_total_cadastros'))
        self.assertContains(response, '2')

    def test_export_views_return_expected_formats_without_full_cpf(self):
        csv_response = self.client.get(reverse('pessoas:export_csv'))
        self.assertEqual(csv_response.status_code, 200)
        self.assertTrue(csv_response['Content-Type'].startswith('text/csv'))
        csv_content = csv_response.content.decode('utf-8')
        self.assertIn('Ana', csv_content)
        self.assertIn(self.pessoa.cpf_mascarado, csv_content)
        self.assertNotIn(self.pessoa.cpf, csv_content)

        json_response = self.client.get(reverse('pessoas:export_json'))
        self.assertEqual(json_response.status_code, 200)
        self.assertEqual(json.loads(json_response.content)[0]['cpf'], self.pessoa.cpf_mascarado)

        xlsx_response = self.client.get(reverse('pessoas:export_xlsx'))
        self.assertEqual(xlsx_response.status_code, 200)
        workbook = load_workbook(BytesIO(xlsx_response.content), read_only=True)
        self.assertEqual(workbook.sheetnames, ['Pessoas'])
        self.assertEqual(workbook.active['A2'].value, 'Ana')
        self.assertEqual(workbook.active['E2'].value, self.pessoa.cpf_mascarado)
        workbook.close()

    def test_export_neutralizes_spreadsheet_formula(self):
        Pessoa.objects.create(
            nome='=1+1',
            sobrenome='Teste',
            data_nascimento=date(2000, 1, 1),
            sexo=Pessoa.SEXO_FEM,
            cpf=cpf_from_base('222333444'),
        )

        csv_content = self.client.get(reverse('pessoas:export_csv')).content.decode('utf-8')
        workbook = load_workbook(
            BytesIO(self.client.get(reverse('pessoas:export_xlsx')).content),
            read_only=True,
        )
        names = [row[0].value for row in workbook.active.iter_rows(min_col=1, max_col=1)]

        self.assertIn("'=1+1", csv_content)
        self.assertIn("'=1+1", names)
        workbook.close()
